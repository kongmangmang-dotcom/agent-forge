"""Document text extraction — port of SolarSense worker TEXT_PARSE_DOCUMENT.

Source reference:
  solarsense-worker/preprocess/multimodal_processors.py::UnstructuredTextParseProcessor

Supports: pdf (PyMuPDF embedded text), docx, html, txt/md and other plain text.
OCR fallback (DeepSeek) is optional via DEEPSEEK_OCR_URL; default is embedded_only/auto
without calling remote OCR when URL is unset.
"""

from __future__ import annotations

import io
import os
import re
from collections import Counter
from dataclasses import dataclass


@dataclass
class ParseResult:
    text: str
    format: str
    char_count: int


def parse_document_bytes(
    data: bytes,
    filename: str,
    *,
    output_format: str = "markdown",
    extract_tables: bool = True,
    ocr_mode: str = "embedded_only",
    ocr_threshold: int = 20,
    page_range: str = "",
) -> ParseResult:
    if not data:
        raise ValueError("文档内容为空")
    if len(data) > 200 * 1024 * 1024:
        raise ValueError("文档过大（超过 200MB）")

    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    if ext == "pdf":
        text = _parse_pdf(
            data,
            extract_tables=extract_tables,
            ocr_mode=ocr_mode,
            ocr_threshold=ocr_threshold,
            page_range=page_range,
        )
    elif ext in {"docx", "doc"}:
        text = _parse_docx(data, extract_tables=extract_tables)
    elif ext in {"html", "htm"}:
        text = _parse_html(data)
    else:
        text = _parse_text(data)

    text = clean_after_parse(text)
    if not text.strip():
        raise ValueError("文档解析结果为空")
    if output_format == "text":
        # strip simple markdown heading markers from docx headings
        text = re.sub(r"(?m)^#{1,6}\s+", "", text)
    return ParseResult(text=text.strip(), format=ext or "txt", char_count=len(text))


def clean_after_parse(text: str) -> str:
    """In-parser clean (UnstructuredTextParseProcessor._clean_text)."""
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    lines = [line.rstrip() for line in text.split("\n")]
    joined = "\n".join(lines)
    if "\n---\n" in joined:
        page_texts = joined.split("\n---\n")
        if len(page_texts) >= 3:
            first_lines = [p.strip().split("\n")[0] for p in page_texts if p.strip()]
            common = Counter(first_lines).most_common(3)
            headers_to_remove = {t for t, c in common if c >= 3 and len(t) < 100}
            if headers_to_remove:
                lines = [ln for ln in lines if ln.strip() not in headers_to_remove]
    return "\n".join(lines).strip()


def _parse_pdf(
    data: bytes,
    *,
    extract_tables: bool,
    ocr_mode: str,
    ocr_threshold: int,
    page_range: str,
) -> str:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("PyMuPDF 未安装，无法解析 PDF") from exc

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        pages = range(len(doc))
        if page_range.strip():
            parts = page_range.split("-")
            start = max(0, int(parts[0]) - 1)
            end = min(len(doc), int(parts[-1])) if len(parts) > 1 else start + 1
            pages = range(start, end)
        page_list = list(pages)

        sections: list[str] = []
        embedded_chars = 0
        for i in page_list:
            page = doc[i]
            text = page.get_text("text") or ""
            if extract_tables:
                try:
                    tables = page.find_tables()
                    for table in tables:
                        df = table.to_pandas()
                        text += "\n" + df.to_markdown(index=False)
                except Exception:
                    pass
            sections.append(text)
            embedded_chars += len(text)

        avg = embedded_chars / max(1, len(page_list))
        ocr_mode = (ocr_mode or "embedded_only").lower()
        use_ocr = ocr_mode == "force_ocr" or (
            ocr_mode == "auto" and avg < ocr_threshold and bool(os.environ.get("DEEPSEEK_OCR_URL"))
        )
        if not use_ocr:
            return "\n\n---\n\n".join(sections)

        # Optional remote OCR — only when explicitly configured
        return _ocr_pdf_pages(doc, page_list)
    finally:
        doc.close()


def _ocr_pdf_pages(doc, page_indices: list[int]) -> str:
    """Best-effort DeepSeek OCR; failures yield empty page slots."""
    import httpx

    deepseek_url = os.environ.get("DEEPSEEK_OCR_URL", "").rstrip("/")
    if not deepseek_url:
        return "\n\n---\n\n".join((doc[i].get_text("text") or "") for i in page_indices)

    import fitz

    dpi = 144
    mat = fitz.Matrix(dpi / 72, dpi / 72)
    sections: list[str] = []
    with httpx.Client(timeout=60.0) as client:
        for page_num in page_indices:
            try:
                pix = doc[page_num].get_pixmap(matrix=mat)
                png = pix.tobytes("png")
                resp = client.post(
                    f"{deepseek_url}/ocr/image",
                    files={"file": (f"page_{page_num}.png", png, "image/png")},
                )
                resp.raise_for_status()
                payload = resp.json()
                text = ""
                if payload.get("success"):
                    text = payload.get("markdown") or payload.get("text") or ""
                sections.append(text)
            except Exception:
                sections.append("")
    return "\n\n---\n\n".join(sections)


def _parse_docx(data: bytes, *, extract_tables: bool) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("python-docx 未安装，无法解析 DOCX") from exc

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style and para.style.name else ""
        if style_name.startswith("Heading"):
            level = style_name[-1] if style_name[-1].isdigit() else "2"
            parts.append(f"{'#' * int(level)} {text}")
        else:
            parts.append(text)
    if extract_tables:
        for table in document.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            header = " | ".join(rows[0])
            sep = " | ".join(["---"] * len(rows[0]))
            body = "\n".join(" | ".join(r) for r in rows[1:])
            parts.append(f"\n{header}\n{sep}\n{body}\n")
    return "\n\n".join(parts)


def _parse_html(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    clean = re.sub(r"<script[^>]*>.*?</script>", "", text, flags=re.DOTALL | re.IGNORECASE)
    clean = re.sub(r"<style[^>]*>.*?</style>", "", clean, flags=re.DOTALL | re.IGNORECASE)
    clean = re.sub(r"<[^>]+>", " ", clean)
    clean = re.sub(r"&\w+;", " ", clean)
    return clean


def _parse_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")
